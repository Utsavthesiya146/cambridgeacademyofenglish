import os
import glob
import re
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def clean_text(text):
    text = re.sub(r'className="[^"]*"', '', text)
    text = re.sub(r'className=\{[^}]*\}', '', text)
    text = re.sub(r'import\s+.*?from\s+[\'"].*?[\'"];?', '', text, flags=re.DOTALL)
    text = re.sub(r'export default function.*?\{', '', text)
    text = re.sub(r'export const metadata.*?\}', '', text, flags=re.DOTALL)
    text = re.sub(r'return\s*\(\s*<>', '', text)
    text = re.sub(r'return\s*\(', '', text)
    text = re.sub(r'\);\s*\}\s*$', '', text)
    text = re.sub(r'\{/\*.*?\*/\}', '', text, flags=re.DOTALL)
    text = re.sub(r'\{.*?\?\s*\((.*?)\)\s*:\s*\((.*?)\)\}', r'\1 \2', text, flags=re.DOTALL)
    text = re.sub(r'\{.*?\}', '', text) # Remove remaining basic jsx expressions
    return text

def extract_text_from_tsx(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # Pre-clean TSX specific stuff before parsing as HTML
    content = clean_text(content)
    
    # Parse as HTML
    soup = BeautifulSoup(content, 'html.parser')
    
    elements = []
    
    for tag in soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'li', 'button', 'a', 'label', 'option', 'th', 'td']):
        text = tag.get_text(separator=' ', strip=True)
        if text and len(text) > 1:
            name = tag.name
            elements.append({'tag': name, 'text': text})
            
    return elements

def main():
    base_dir = r"F:\Cambridge Academy of English\CambridgeAcademyNew\src\app"
    page_files = glob.glob(os.path.join(base_dir, "**", "page.tsx"), recursive=True)
    
    doc = Document()
    
    # Title Page
    title = doc.add_heading('Cambridge Academy of English', level=0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    subtitle = doc.add_paragraph('Complete Website Content\nContent Reference Document')
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_page_break()
    
    # Table of Contents placeholder
    doc.add_heading('Table of Contents', level=1)
    doc.add_paragraph('Pages documented in this file:')
    for pf in page_files:
        route = pf.replace(base_dir, '').replace('\\page.tsx', '').replace('\\', '/')
        if not route: route = '/'
        doc.add_paragraph(f"- {route}")
    doc.add_page_break()
    
    total_pages = len(page_files)
    total_sections = 0
    empty_pages = 0

    for pf in page_files:
        route = pf.replace(base_dir, '').replace('\\page.tsx', '').replace('\\', '/')
        if not route: route = '/'
        
        doc.add_heading(f'Route: {route}', level=1)
        
        elements = extract_text_from_tsx(pf)
        
        if not elements:
            doc.add_paragraph("Content currently not implemented in the project.")
            empty_pages += 1
        else:
            for el in elements:
                tag = el['tag']
                text = el['text']
                if text.startswith("import") or text.startswith("export"):
                    continue
                if tag == 'h1':
                    doc.add_heading(text, level=1)
                    total_sections += 1
                elif tag == 'h2':
                    doc.add_heading(text, level=2)
                    total_sections += 1
                elif tag == 'h3' or tag == 'h4':
                    doc.add_heading(text, level=3)
                elif tag == 'li':
                    doc.add_paragraph(text, style='List Bullet')
                elif tag in ['th', 'td', 'option', 'label', 'button']:
                    doc.add_paragraph(f"[{tag.upper()}]: {text}")
                else:
                    doc.add_paragraph(text)
                    
        doc.add_page_break()
        
    doc_path = r"F:\Cambridge Academy of English\CambridgeAcademyNew\Cambridge_Academy_of_English_Complete_Website_Content.docx"
    doc.save(doc_path)
    print(f"DONE. Total Pages: {total_pages}. Total Sections: {total_sections}. Empty Pages: {empty_pages}.")
    print(f"Path: {doc_path}")

if __name__ == '__main__':
    main()
